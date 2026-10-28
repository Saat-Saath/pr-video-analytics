"""Build PR_Video_Analytics_Master.xlsx from analytics_data.json (output of merge_results.py).
Replaces the Codex-only build_master.mjs. Usage: python build_workbook.py
"""
import json
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule

ROOT = Path(__file__).resolve().parent
DATA = json.loads((ROOT / "analytics_data.json").read_text(encoding="utf8"))
try:  # cost per video from the PR Video Costs sheet (empty when no amounts are filled in yet)
    from costs import allocate
    COST, COST_REPORT, PAYEE = allocate(DATA["master"])
except (SystemExit, Exception) as _e:
    print("costs skipped:", _e)
    COST, COST_REPORT, PAYEE = {}, [], {}
OUT = ROOT.parent / "outputs" / "PR_Video_Analytics_Master.xlsx"

# Names shown in the workbook (the internal key stays the same so videos keep matching across refreshes).
DISPLAY_NAMES = {"Kalpana - Disease Channels": "Inhouse ai videos"}
for _rec in DATA["master"] + DATA["occurrences"] + DATA["config"]:
    _rec["source_group"] = DISPLAY_NAMES.get(_rec.get("source_group"), _rec.get("source_group"))

NAVY, LIGHT, GRAY, DARK = "1F4E78", "D9EAF7", "F2F2F2", "202124"
F = "Arial"
thin = Side(style="thin", color="D9D9D9")
HEAD = dict(font=Font(name=F, bold=True, color="FFFFFF", size=10), fill=PatternFill("solid", fgColor=NAVY),
            alignment=Alignment(horizontal="center", vertical="center", wrap_text=True))


def style_header(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.alignment = HEAD["font"], HEAD["fill"], HEAD["alignment"]


def d(v):
    if not v: return None
    try: return datetime.fromisoformat(str(v)[:10]).date()
    except ValueError: return None


def dt(v):
    if not v: return None
    try: return datetime.fromisoformat(str(v).replace("Z", "+00:00")).replace(tzinfo=None, microsecond=0)
    except ValueError: return None


def add_table(ws, name, ncols, nrows, style="TableStyleLight9"):
    ref = f"A1:{get_column_letter(ncols)}{max(nrows + 1, 2)}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()
dash = wb.active; dash.title = "Dashboard"
master_ws = wb.create_sheet("Master Content")
src_ws = wb.create_sheet("Source Rows")
snap_ws = wb.create_sheet("Analytics Snapshots")
setup = wb.create_sheet("Setup")
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False

# ---------- Master Content ----------
master = DATA["master"]
mcols = [("Master ID", 11), ("Platform", 12), ("Content type", 13), ("Source group", 24), ("Channel / influencer", 24),
         ("Account", 24), ("Post URL", 46), ("Published", 12),
         ("Day 2 views", 12), ("Day 2 likes", 11), ("Day 2 status", 20),
         ("Day 7 views", 12), ("Day 7 likes", 11), ("Day 7 status", 20),
         ("Day 30 views", 12), ("Day 30 likes", 11), ("Day 30 status", 20),
         ("All-time views", 13), ("All-time likes", 12), ("All-time comments", 12), ("All-time shares", 12), ("Engagement rate", 12),
         ("Last fetched", 18), ("Data source", 12), ("Fetch status", 15), ("Sheet status", 14), ("Source rows", 22), ("Occurrences", 11),
         ("Cost (₹)", 11), ("₹ per 1,000 views", 12), ("Paid to", 22)]
for i, (h, _) in enumerate(mcols, 1):
    master_ws.cell(row=1, column=i, value=h)
style_header(master_ws, 1, 1, len(mcols))
widths(master_ws, [w for _, w in mcols])
for r, m in enumerate(master, start=2):
    row = [m["master_id"], m["platform"], m["content_type"], m["source_group"], m.get("channel") or m.get("influencer") or "",
           m.get("account") or "", m["canonical_url"], d(m.get("published_date")),
           m.get("d2_views"), m.get("d2_likes"), m.get("d2_status"),
           m.get("d7_views"), m.get("d7_likes"), m.get("d7_status"),
           m.get("d30_views"), m.get("d30_likes"), m.get("d30_status"),
           m.get("latest_views"), m.get("latest_likes"), m.get("latest_comments"), m.get("latest_shares"), m.get("latest_engagement_rate"),
           dt(m.get("last_analytics_at")), m.get("data_source") or "", m.get("fetch_status"), m.get("sheet_status") or "In sheet",
           m.get("source_rows"), m.get("occurrence_count"), COST.get(m["master_id"]), None, PAYEE.get(m["master_id"])]
    for c, v in enumerate(row, start=1):
        cell = master_ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=F, size=10, color=DARK)
    master_ws.cell(row=r, column=30, value=f'=IF(AND(ISNUMBER(AC{r}),N(R{r})>0),AC{r}/R{r}*1000,"")')
    master_ws.cell(row=r, column=7).hyperlink = m["canonical_url"]
    master_ws.cell(row=r, column=7).font = Font(name=F, size=10, color="1155CC", underline="single")
mend = len(master) + 1
for c in master_ws[f"H2:H{mend}"]: c[0].number_format = "dd-mmm-yyyy"
for col in "IJLMOPRSTU":
    for c in master_ws[f"{col}2:{col}{mend}"]: c[0].number_format = "#,##0"
for c in master_ws[f"V2:V{mend}"]: c[0].number_format = "0.0%"
for col in ("AC", "AD"):
    for c in master_ws[f"{col}2:{col}{mend}"]: c[0].number_format = '"₹"#,##0.00'
for c in master_ws[f"W2:W{mend}"]: c[0].number_format = "dd-mmm-yyyy hh:mm"
add_table(master_ws, "MasterContent", len(mcols), len(master))
master_ws.freeze_panes = "C2"
amber = PatternFill("solid", fgColor="FFF2CC"); red = PatternFill("solid", fgColor="FCE8E6"); green = PatternFill("solid", fgColor="E2F0D9")
rng = f"Y2:Y{mend}"
master_ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=green))
master_ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"No view count"'], fill=amber))
master_ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"No public link"'], fill=amber))
master_ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Fetch failed"'], fill=red))
master_ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Not fetched"'], fill=red))

# ---------- Source Rows ----------
occ = DATA["occurrences"]
scols = [("Source group", 24), ("Workbook", 30), ("Tab", 16), ("Row", 7), ("Cell", 7), ("Date text", 14), ("Published", 12),
         ("Influencer", 18), ("Channel", 24), ("Profile URL", 30), ("Declared platform", 14), ("Posting status", 16),
         ("Original URL", 46), ("Canonical URL", 46), ("Detected platform", 14)]
keys = ["source_group", "source_workbook", "source_tab", "source_row", "source_cell", "source_date_text", "published_date",
        "influencer", "channel", "profile_url", "declared_platform", "posting_status", "url", "canonical_url", "detected_platform"]
for i, (h, _) in enumerate(scols, 1): src_ws.cell(row=1, column=i, value=h)
style_header(src_ws, 1, 1, len(scols)); widths(src_ws, [w for _, w in scols])
for r, o in enumerate(occ, start=2):
    for c, k in enumerate(keys, start=1):
        v = o.get(k)
        if k == "detected_platform": v = "Google Drive Asset" if "drive.google" in (o.get("canonical_url") or "") else ("YouTube" if str(v).lower() == "youtube" else v)
        src_ws.cell(row=r, column=c, value=v).font = Font(name=F, size=10)
add_table(src_ws, "SourceRows", len(scols), len(occ)); src_ws.freeze_panes = "A2"

# ---------- Analytics Snapshots ----------
snaps = DATA["snapshots"]
ncols = [("Snapshot ID", 12), ("Master ID", 11), ("Snapshot type", 13), ("Captured at", 18), ("Age (days)", 10), ("Views", 11),
         ("Likes", 10), ("Comments", 10), ("Shares", 10), ("Engagement rate", 12), ("Account", 26), ("Data source", 12),
         ("Fetch status", 12), ("Error detail", 40)]
for i, (h, _) in enumerate(ncols, 1): snap_ws.cell(row=1, column=i, value=h)
style_header(snap_ws, 1, 1, len(ncols)); widths(snap_ws, [w for _, w in ncols])
for r, s in enumerate(snaps, start=2):
    vals = [s["snapshot_id"], s["master_id"], s["snapshot_type"], dt(s["captured_at"]), s["age_days"], s["views"], s["likes"],
            s["comments"], s["shares"], None, s.get("account"), s["source"], s["fetch_status"], s.get("error") or ""]
    for c, v in enumerate(vals, start=1): snap_ws.cell(row=r, column=c, value=v).font = Font(name=F, size=10)
    snap_ws.cell(row=r, column=10, value=f'=IF(AND(ISNUMBER(F{r}),F{r}>0),(N(G{r})+N(H{r})+N(I{r}))/F{r},"")').number_format = "0.0%"
    snap_ws.cell(row=r, column=4).number_format = "dd-mmm-yyyy hh:mm"
    for col in (6, 7, 8, 9): snap_ws.cell(row=r, column=col).number_format = "#,##0"
add_table(snap_ws, "Snapshots", len(ncols), len(snaps)); snap_ws.freeze_panes = "A2"

# ---------- Dashboard (live formulas over Master Content) ----------
M = "'Master Content'"
def mc(col): return f"{M}!${col}$2:${col}${mend}"
dash.column_dimensions["A"].width = 30
for col in "BCDEFGH": dash.column_dimensions[col].width = 15
dash["A1"] = "PR Video Analytics Master"; dash["A1"].font = Font(name=F, size=16, bold=True, color=NAVY)
dash["A2"] = f"Inhouse ai videos, Sanjay Agency and Internal PR · data refreshed {datetime.now():%d %b %Y %H:%M}"
dash["A2"].font = Font(name=F, size=10, italic=True, color="5F6368")

kpis = [("Videos tracked", f"=COUNTA({mc('A')})", "#,##0"),
        ("Total views", f"=SUM({mc('R')})", "#,##0"),
        ("Likes + comments", f"=SUM({mc('S')})+SUM({mc('T')})", "#,##0"),
        ("Engagement rate", f"=IFERROR((SUMIFS({mc('S')},{mc('R')},\">0\")+SUMIFS({mc('T')},{mc('R')},\">0\")+SUMIFS({mc('U')},{mc('R')},\">0\"))/SUM({mc('R')}),0)", "0.0%"),
        ("Missing data", f"=COUNTA({mc('A')})-COUNTIF({mc('Y')},\"OK\")", "#,##0")]
for i, (label, f, fmt) in enumerate(kpis):
    col = get_column_letter(1 + i) if i == 0 else get_column_letter(1 + i)
    lab = dash.cell(row=4, column=1 + i, value=label); val = dash.cell(row=5, column=1 + i, value=f)
    lab.font = Font(name=F, size=9, color="5F6368"); val.font = Font(name=F, size=16, bold=True, color=NAVY); val.number_format = fmt
    for c in (lab, val): c.fill = PatternFill("solid", fgColor=GRAY); c.alignment = Alignment(horizontal="left")

def block(top, title, header, rows):
    dash.cell(row=top, column=1, value=title).font = Font(name=F, size=11, bold=True, color=NAVY)
    for i, h in enumerate(header, 1): dash.cell(row=top + 1, column=i, value=h)
    style_header(dash, top + 1, 1, len(header))
    for r, vals in enumerate(rows, start=top + 2):
        for c, v in enumerate(vals, start=1):
            cell = dash.cell(row=r, column=c, value=v); cell.font = Font(name=F, size=10)
            cell.border = Border(bottom=thin)
    return top + 2 + len(rows) + 1

hdr = ["", "Videos", "With views", "Views", "Likes", "Comments", "Shares", "Eng. rate"]
def agg_row(label, col):
    crit = f'{mc(col)},"{label}"'
    r = [label, f"=COUNTIFS({crit})", f'=COUNTIFS({crit},{mc("R")},">=0")', f"=SUMIFS({mc('R')},{crit})",
         f"=SUMIFS({mc('S')},{crit})", f"=SUMIFS({mc('T')},{crit})", f"=SUMIFS({mc('U')},{crit})"]
    r.append(f"=IFERROR((SUMIFS({mc('S')},{crit},{mc('R')},\">0\")+SUMIFS({mc('T')},{crit},{mc('R')},\">0\")+SUMIFS({mc('U')},{crit},{mc('R')},\">0\"))/SUMIFS({mc('R')},{crit}),\"\")")
    return r
nxt = block(7, "By platform", ["Platform"] + hdr[1:], [agg_row(p, "B") for p in ("Facebook", "Instagram", "YouTube", "Google Drive Asset")])
start_src = nxt
nxt = block(nxt, "By source", ["Source group"] + hdr[1:], [agg_row(s["source_group"], "D") for s in DATA["config"]])
for r in range(8, nxt):
    for c in range(2, 8): dash.cell(row=r, column=c).number_format = "#,##0"
    dash.cell(row=r, column=8).number_format = "0.0%"

top10 = sorted([m for m in master if m.get("latest_views")], key=lambda m: -m["latest_views"])[:10]
t = nxt
dash.cell(row=t, column=1, value="Top 10 videos by views (at last refresh)").font = Font(name=F, size=11, bold=True, color=NAVY)
for i, h in enumerate(["Video", "Platform", "Source", "Views", "Likes", "Published", "Account"], 1): dash.cell(row=t + 1, column=i, value=h)
style_header(dash, t + 1, 1, 7)
for r, m in enumerate(top10, start=t + 2):
    vals = [m["master_id"], m["platform"], m["source_group"], m["latest_views"], m["latest_likes"], d(m.get("published_date")), m.get("account")]
    for c, v in enumerate(vals, 1):
        cell = dash.cell(row=r, column=c, value=v); cell.font = Font(name=F, size=10); cell.border = Border(bottom=thin)
    dash.cell(row=r, column=1).hyperlink = m["canonical_url"]; dash.cell(row=r, column=1).font = Font(name=F, size=10, color="1155CC", underline="single")
    dash.cell(row=r, column=4).number_format = "#,##0"; dash.cell(row=r, column=5).number_format = "#,##0"; dash.cell(row=r, column=6).number_format = "dd-mmm-yyyy"
t = t + 2 + len(top10) + 1
dash.cell(row=t, column=1, value="Milestones (Day 2 / Day 7 / Day 30 after upload)").font = Font(name=F, size=11, bold=True, color=NAVY)
for i, h in enumerate(["Milestone", "Captured", "Captured late", "Due now", "Due later", "Missed", "Views captured"], 1): dash.cell(row=t + 1, column=i, value=h)
style_header(dash, t + 1, 1, 7)
for r, (lab, colv, cols) in enumerate([("Day 2", "I", "K"), ("Day 7", "L", "N"), ("Day 30", "O", "Q")], start=t + 2):
    vals = [lab, f'=COUNTIF({mc(cols)},"Captured")', f'=COUNTIF({mc(cols)},"Captured late*")', f'=COUNTIF({mc(cols)},"Due now")',
            f'=COUNTIF({mc(cols)},"Due 2*")', f'=COUNTIF({mc(cols)},"Missed")', f"=SUM({mc(colv)})"]
    for c, v in enumerate(vals, 1):
        cell = dash.cell(row=r, column=c, value=v); cell.font = Font(name=F, size=10); cell.border = Border(bottom=thin)
        if c > 1: cell.number_format = "#,##0"
t = t + 2 + 3 + 1
fix = [("Fetched, but no view count", f'=COUNTIF({mc("Y")},"No view count")', "Facebook hides some counts; re-check later"),
       ("Fetch failed", f'=COUNTIF({mc("Y")},"Fetch failed")', "Link deleted/private or scraper error; open and verify"),
       ("Drive asset, no public post link", f'=COUNTIF({mc("Y")},"No public link")', "Ask Internal PR for the live post URL"),
       ("Missing publish date", f'=COUNTBLANK({mc("H")})', "7/30-day dates can't be scheduled")]
nxt = block(t, "To fix", ["Issue", "Count", "Next step"], fix)

# ---------- Cost per view ----------
cost_ws = wb.create_sheet("Cost per view", index=1)
cost_ws.sheet_view.showGridLines = False
cost_ws["A1"] = "Cost per view"; cost_ws["A1"].font = Font(name=F, size=14, bold=True, color=NAVY)
cost_ws["A2"] = ("From the PR Video Costs sheet. A monthly amount is shared equally by that account's videos posted that month. "
                 "₹ per 1,000 views uses all-time views and counts only videos that have a view count.")
cost_ws["A2"].font = Font(name=F, size=10, italic=True, color="5F6368")
for col, w in zip("ABCDEFGH", (30, 22, 10, 14, 14, 14, 16, 14)): cost_ws.column_dimensions[col].width = w


def cost_groups(key):
    g = {}
    for m in master:
        c = COST.get(m["master_id"])
        if not c: continue
        k = key(m); e = g.setdefault(k, {"source": m["source_group"], "n": 0, "spend": 0.0, "cspend": 0.0, "views": 0})
        e["n"] += 1; e["spend"] += c
        if m.get("latest_views") is not None: e["cspend"] += c; e["views"] += m["latest_views"]
    out = [[k, e["source"], e["n"], round(e["spend"]), e["views"], round(e["spend"] / e["n"]),
            (e["cspend"] / e["views"] * 1000) if e["views"] else None] for k, e in g.items()]
    return sorted(out, key=lambda r: (r[6] is None, r[6] or 0))


hdr_c = ["Paid to", "Source", "Videos", "Spend (₹)", "Views", "Per video (₹)", "₹ per 1,000 views"]
row_i = 4
for title, key in (("By source", lambda m: m["source_group"]),
                   ("By account (cheapest views first)", lambda m: PAYEE.get(m["master_id"]) or m["source_group"])):
    cost_ws.cell(row=row_i, column=1, value=title).font = Font(name=F, size=11, bold=True, color=NAVY)
    for i, h in enumerate(hdr_c, 1): cost_ws.cell(row=row_i + 1, column=i, value=h)
    style_header(cost_ws, row_i + 1, 1, len(hdr_c))
    data_rows = cost_groups(key)
    if not data_rows:
        cost_ws.cell(row=row_i + 2, column=1, value="No amounts yet: fill in the PR Video Costs sheet").font = Font(name=F, size=10, italic=True)
    for r, vals in enumerate(data_rows, start=row_i + 2):
        for c, v in enumerate(vals, 1):
            cell = cost_ws.cell(row=r, column=c, value=v); cell.font = Font(name=F, size=10); cell.border = Border(bottom=thin)
        for c in (4, 6): cost_ws.cell(row=r, column=c).number_format = '"₹"#,##0'
        cost_ws.cell(row=r, column=5).number_format = "#,##0"; cost_ws.cell(row=r, column=7).number_format = '"₹"#,##0.00'
    row_i += 2 + max(1, len(data_rows)) + 1
cost_ws.cell(row=row_i, column=1, value="Rows read from the PR Video Costs sheet").font = Font(name=F, size=11, bold=True, color=NAVY)
for i, h in enumerate(["Sheet row", "Source", "Account", "Month", "Amount (₹)", "Basis", "Videos matched / problem", "Promised / month"], 1):
    cost_ws.cell(row=row_i + 1, column=i, value=h)
style_header(cost_ws, row_i + 1, 1, 8)
for r, c in enumerate(COST_REPORT, start=row_i + 2):
    vals = [c["row"], c["source"], c["account"] or "(all)", c["month"] or "(all months)", c["amount"], c["basis"], c["problem"] or c["videos"], c.get("promised")]
    for k, v in enumerate(vals, 1):
        cost_ws.cell(row=r, column=k, value=v).font = Font(name=F, size=10, color="B3261E" if c["problem"] and k == 7 else DARK)
    cost_ws.cell(row=r, column=5).number_format = '"₹"#,##0'

# ---------- Setup ----------
setup.column_dimensions["A"].width = 26; setup.column_dimensions["B"].width = 36
for col in "CDEF": setup.column_dimensions[col].width = 30
setup["A1"] = "Sources and refresh"; setup["A1"].font = Font(name=F, size=14, bold=True, color=NAVY)
for i, h in enumerate(["Source group", "Spreadsheet", "Tabs", "Parser", "Source URL"], 1): setup.cell(row=3, column=i, value=h)
style_header(setup, 3, 1, 5)
for r, c in enumerate(DATA["config"], start=4):
    for i, k in enumerate(["source_group", "spreadsheet_title", "tabs", "parser", "source_url"], 1):
        setup.cell(row=r, column=i, value=c[k]).font = Font(name=F, size=10)
steps = [("Instagram metrics", "Apify Instagram Scraper (apify_instagram.mjs) - plays, likes, comments, date"),
         ("Facebook / YouTube metrics", "Bright Data Scraper API (brightdata_api_pipeline.mjs)"),
         ("Daily refresh", "python refresh.py start, then python refresh.py finish (syncs sheets, fetches due milestones, rebuilds)"),
         ("Rebuild workbook", "python analytics_master_build/build_workbook.py"),
         ("Milestones", "Day 2 = captured 2-6 days after upload, Day 7 = 7-29 days, Day 30 = 30+ days; late captures are flagged"),
         ("Keys", "BRIGHTDATA_API_KEY and APIFY_TOKEN in .env (never commit)")]
setup.cell(row=9, column=1, value="Refresh workflow").font = Font(name=F, size=11, bold=True, color=NAVY)
for r, (a, b) in enumerate(steps, start=10):
    setup.cell(row=r, column=1, value=a).font = Font(name=F, size=10, bold=True)
    setup.cell(row=r, column=2, value=b).font = Font(name=F, size=10)

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("saved", OUT)
